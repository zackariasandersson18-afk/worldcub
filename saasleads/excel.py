"""Skriver leads till Excel och bygger Hitta.se-söklänkar."""
from __future__ import annotations

from datetime import date
from urllib.parse import quote_plus

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import Lead, format_orgnr

HITTA_SEARCH = "https://www.hitta.se/s%C3%B6k?vad="


def hitta_link(first: str, last: str, city: str) -> str:
    """Söklänk för manuell uppslagning. Hitta.se anropas aldrig av verktyget."""
    query = " ".join(p for p in (first, last, city) if p).strip()
    return HITTA_SEARCH + quote_plus(query) if first and last else ""


COLUMNS = [
    ("Företagsnamn", 32, lambda l: l.name),
    ("Organisationsnummer", 16, lambda l: format_orgnr(l.orgnr)),
    ("Hemsida", 30, lambda l: l.website),
    ("Antal anställda", 10, lambda l: l.employees),
    ("Storleksklass (SCB)", 16, lambda l: l.size_class),
    ("Ort", 16, lambda l: l.city),
    ("VD/grundare förnamn", 16, lambda l: l.ceo_first_name),
    ("VD/grundare efternamn", 18, lambda l: l.ceo_last_name),
    ("Roll", 22, lambda l: l.ceo_role),
    ("Hitta.se-sökning", 20,
     lambda l: hitta_link(l.ceo_first_name, l.ceo_last_name, l.city)),
    ("SaaS-poäng", 9, lambda l: l.saas_score),
    ("SaaS-signaler", 40, lambda l: ", ".join(l.saas_keywords)),
    ("SNI", 14, lambda l: ", ".join(l.sni_codes)),
    ("Hemsida verifierad via", 12, lambda l: l.website_confidence),
    ("Anställda enligt", 22, lambda l: l.employees_source),
    ("Källor", 22, lambda l: ", ".join(l.sources)),
    ("Anledning/anteckning", 45,
     lambda l: "; ".join(x for x in [l.reason, *l.notes] if x)),
]
LINK_COLUMNS = {"Hemsida", "Hitta.se-sökning"}
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")


def _sheet(wb: Workbook, title: str, leads: list[Lead], first: bool = False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    for c, (header, width, _) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=c, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        ws.column_dimensions[get_column_letter(c)].width = width
    for r, lead in enumerate(leads, start=2):
        for c, (header, _, getter) in enumerate(COLUMNS, start=1):
            value = getter(lead)
            cell = ws.cell(row=r, column=c, value=value if value != "" else None)
            if header in LINK_COLUMNS and value:
                cell.hyperlink = value
                cell.style = "Hyperlink"
                if header == "Hitta.se-sökning":
                    cell.value = "Sök på Hitta"
    ws.freeze_panes = "B2"
    if leads:
        ws.auto_filter.ref = ws.dimensions


def write_workbook(path: str, leads: list[Lead], meta: dict | None = None) -> None:
    wb = Workbook()
    by_status = lambda s: [l for l in leads if l.status == s]
    _sheet(wb, "Leads", by_status("lead"), first=True)
    _sheet(wb, "Att granska", by_status("granska"))
    _sheet(wb, "Bortfiltrerade", by_status("filtrerad"))

    about = wb.create_sheet("Om")
    rows = [("Skapad", date.today().isoformat()), *(meta or {}).items(),
            ("Obs", "VD-namn är personuppgifter (GDPR). Använd för B2B-kontakt "
                    "med berättigat intresse, informera vid första kontakt och "
                    "respektera NIX/avregistrering.")]
    for r, (k, v) in enumerate(rows, start=1):
        about.cell(row=r, column=1, value=k).font = Font(bold=True)
        about.cell(row=r, column=2, value=str(v))
    about.column_dimensions["A"].width = 22
    about.column_dimensions["B"].width = 100
    wb.save(path)
