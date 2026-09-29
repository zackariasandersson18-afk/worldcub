import io
import zipfile

import pytest
from openpyxl import load_workbook

from saasleads import saas
from saasleads.excel import hitta_link, write_workbook
from saasleads.ixbrl import Signer, choose_ceo, extract_xhtml_from_zip, parse_report
from saasleads.models import Lead, format_orgnr, normalize_orgnr
from saasleads.pipeline import Pipeline
from saasleads.polite import BlockedError, PoliteClient, is_blocked_domain
from saasleads.scb import row_to_lead
from saasleads.website import candidate_domains, verify

# ---------------------------------------------------------------- polite ---


class FakeResp:
    def __init__(self, status=200, text="", url="", headers=None):
        self.status_code, self.text, self.url = status, text, url
        self.headers = headers or {"Content-Type": "text/html"}


class FakeSession:
    def __init__(self, pages):
        self.pages, self.calls, self.headers = pages, [], {}

    def get(self, url, **kw):
        return self.request("GET", url, **kw)

    def request(self, method, url, **kw):
        self.calls.append(url)
        status, text = self.pages.get(url, (404, ""))
        return FakeResp(status, text, url)


def make_client(pages, **kw):
    t = {"now": 0.0, "slept": []}

    def sleep(s):
        t["slept"].append(s)
        t["now"] += s
    client = PoliteClient(FakeSession(pages), sleep=sleep,
                          clock=lambda: t["now"], **kw)
    return client, t


@pytest.mark.parametrize("url", [
    "https://www.hitta.se/sök?vad=x", "https://merinfo.se/", "https://www.ratsit.se",
    "https://www.allabolag.se/5560000000"])
def test_blocked_domains_are_never_requested(url):
    client, _ = make_client({})
    assert is_blocked_domain(url)
    with pytest.raises(BlockedError):
        client.get_page(url)
    assert client.session.calls == []


def test_robots_disallow_is_respected():
    client, _ = make_client({
        "https://ex.se/robots.txt": (200, "User-agent: *\nDisallow: /"),
        "https://ex.se/": (200, "<html></html>")})
    with pytest.raises(BlockedError):
        client.get_page("https://ex.se/")
    assert client.session.calls == ["https://ex.se/robots.txt"]


def test_missing_robots_allows_and_delay_is_applied():
    client, t = make_client({"https://ex.se/": (200, "ok")}, min_delay=3.0)
    assert client.get_page("https://ex.se/").text == "ok"
    assert client.get_page("https://ex.se/").text == "ok"
    # robots + två sidor på samma domän => två väntor à 3 s
    assert t["slept"] == [3.0, 3.0]


def test_crawl_delay_from_robots():
    client, t = make_client({
        "https://ex.se/robots.txt": (200, "User-agent: *\nCrawl-delay: 10"),
        "https://ex.se/": (200, "ok")}, min_delay=1.0)
    client.get_page("https://ex.se/")
    assert t["slept"] == [10.0]


def test_robots_server_error_skips_site():
    client, _ = make_client({"https://ex.se/robots.txt": (503, "")})
    assert not client.can_fetch("https://ex.se/")

# ----------------------------------------------------------------- ixbrl ---


XHTML = """<html xmlns:ix="http://www.xbrl.org/2013/inlineXBRL"><body>
<ix:nonFraction name="se-gen-base:MedelantaletAnstallda" contextRef="period1">9</ix:nonFraction>
<ix:nonFraction name="se-gen-base:MedelantaletAnstallda" contextRef="period0">12</ix:nonFraction>
<ix:nonNumeric name="se-gen-base:UnderskriftHandlingTilltalsnamn" tupleRef="t1">Karin</ix:nonNumeric>
<ix:nonNumeric name="se-gen-base:UnderskriftHandlingEfternamn" tupleRef="t1">Lund</ix:nonNumeric>
<ix:nonNumeric name="se-gen-base:UnderskriftHandlingRoll" tupleRef="t1">Styrelseordförande</ix:nonNumeric>
<ix:nonNumeric name="se-gen-base:UnderskriftHandlingTilltalsnamn" tupleRef="t2">Erik</ix:nonNumeric>
<ix:nonNumeric name="se-gen-base:UnderskriftHandlingEfternamn" tupleRef="t2">Öberg</ix:nonNumeric>
<ix:nonNumeric name="se-gen-base:UnderskriftHandlingRoll" tupleRef="t2">Verkställande direktör</ix:nonNumeric>
<ix:nonNumeric name="se-bol-base:UnderskriftFaststallelseintygForetradareTilltalsnamn" tupleRef="f1">Karin</ix:nonNumeric>
</body></html>"""


def test_parse_report_employees_and_ceo():
    report = parse_report(XHTML)
    assert report.employees == 12
    assert [s.full_name for s in report.signers] == ["Karin Lund", "Erik Öberg"]
    ceo = choose_ceo(report.signers)
    assert (ceo.first_name, ceo.last_name, ceo.role) == (
        "Erik", "Öberg", "Verkställande direktör")


def test_choose_ceo_falls_back_to_chair_then_first():
    assert choose_ceo([Signer("A", "B", "Ledamot"),
                       Signer("C", "D", "Ordförande")]).first_name == "C"
    assert choose_ceo([Signer("A", "B", "Styrelseledamot")]).first_name == "A"
    assert choose_ceo([]) is None


def test_text_fallback_and_zip():
    html = "<html><body><p>Stockholm 2025-03-01</p><p>Sara Nilsson</p>" \
           "<p>Verkställande direktör</p></body></html>"
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("arsredovisning.xhtml", html)
    report = parse_report(extract_xhtml_from_zip(buf.getvalue()))
    assert choose_ceo(report.signers).full_name == "Sara Nilsson"

# ------------------------------------------------------------- saas/web ---


SAAS_HTML = """<html><head><title>Planly</title></head><body>
<a href="/pricing">Priser</a> Plattformen för schemaläggning.
Prova gratis i 14 dagar. Från 99 kr/mån per användare. Logga in</body></html>"""
CONSULT_HTML = """<html><body>Vi är IT-konsulter som hjälper er med
systemutveckling. Våra konsulter arbetar på timpris.</body></html>"""


def test_saas_scoring():
    score, hits, ok = saas.assess([saas.page_text(SAAS_HTML)])
    assert ok and "gratis provperiod" in hits and "pricing" in hits
    score, hits, ok = saas.assess([saas.page_text(CONSULT_HTML)])
    assert not ok and "-konsult" in hits


def test_candidate_domains_strip_legal_suffix():
    doms = candidate_domains("Planly Software AB")
    assert doms[:2] == ["planlysoftware.se", "planlysoftware.com"]
    assert "planly.se" in doms
    assert candidate_domains("Åre Molntjänster AB (publ)")[0] == "aremolntjanster.se"


def test_verify_website():
    assert verify("<p>Org.nr 556123-4567</p>", "X AB", "5561234567") == "orgnr"
    assert verify("<title>Planly</title>", "Planly AB", "5561234567") == "namn"
    assert verify("<p>This domain is for sale</p>", "Planly AB", "1") == ""

# --------------------------------------------------------- models/scb/excel -


def test_orgnr_normalization():
    assert normalize_orgnr("16556123-4567") == "5561234567"
    assert format_orgnr("5561234567") == "556123-4567"


def test_scb_row_to_lead_fuzzy_keys():
    lead = row_to_lead({"PeOrgNr": "165561234567", "Företagsnamn": "Planly AB",
                        "PostOrt": "MALMÖ", "Bransch_1, kod": "62010",
                        "Storleksklass": "5-9 anställda"})
    assert (lead.orgnr, lead.city, lead.sni_codes) == ("5561234567", "MALMÖ", ["62010"])
    assert row_to_lead({"Företagsnamn": "Utan orgnr"}) is None


def test_hitta_link_is_only_a_search_url():
    link = hitta_link("Erik", "Öberg", "Malmö")
    assert link == "https://www.hitta.se/s%C3%B6k?vad=Erik+%C3%96berg+Malm%C3%B6"
    assert hitta_link("", "", "Malmö") == ""


def test_workbook(tmp_path):
    leads = [
        Lead("5561234567", "Planly AB", city="Malmö", ceo_first_name="Erik",
             ceo_last_name="Öberg", website="https://planly.se", status="lead"),
        Lead("5560000001", "Konsult AB", status="filtrerad", reason="konsult"),
    ]
    path = tmp_path / "leads.xlsx"
    write_workbook(str(path), leads, {"SNI": "62010"})
    wb = load_workbook(path)
    assert wb.sheetnames == ["Leads", "Att granska", "Bortfiltrerade", "Om"]
    ws = wb["Leads"]
    headers = [c.value for c in ws[1]]
    row = dict(zip(headers, ws[2]))
    assert row["Organisationsnummer"].value == "556123-4567"
    assert row["Hitta.se-sökning"].hyperlink.target.startswith("https://www.hitta.se/")
    assert wb["Bortfiltrerade"].max_row == 2

# -------------------------------------------------------------- pipeline ---


class FakeBV:
    def __init__(self, employees):
        self.employees = employees

    def organisation(self, orgnr):
        return {"name": "", "city": "Malmö", "sni_codes": ["62010"],
                "business_description": ""}

    def latest_annual_report(self, orgnr):
        return parse_report(XHTML.replace(">12<", f">{self.employees}<"))


class FakeWeb:
    def __init__(self, html):
        self.html = html

    def find(self, name, orgnr, known):
        return ("https://planly.se/", "orgnr", self.html) if self.html else ("", "", "")

    def pricing_page(self, url, html):
        return ""


def test_pipeline_statuses(tmp_path):
    make = lambda: Lead("5561234567", "Planly AB", sources=["SCB"])
    lead = Pipeline(FakeBV(12), FakeWeb(SAAS_HTML)).process(make())
    assert lead.status == "lead" and lead.ceo_last_name == "Öberg"
    assert lead.employees == 12 and lead.website == "https://planly.se/"

    assert Pipeline(FakeBV(120), FakeWeb(SAAS_HTML)).process(make()).status == "filtrerad"
    assert Pipeline(FakeBV(1), FakeWeb(SAAS_HTML)).process(make()).status == "filtrerad"
    assert Pipeline(FakeBV(12), FakeWeb(CONSULT_HTML)).process(make()).status == "filtrerad"
    assert Pipeline(FakeBV(12), FakeWeb("")).process(make()).status == "granska"

    # cachen återanvänds vid nästa körning
    cached = Pipeline(FakeBV(12), FakeWeb(SAAS_HTML), cache_dir=str(tmp_path))
    cached.process(make())
    again = Pipeline(None, None, cache_dir=str(tmp_path)).process(make())
    assert again.status == "lead" and again.ceo_first_name == "Erik"
